"""Job-offer parsing, cleaning and heuristic fallbacks."""

import io
import re
from pathlib import Path

MAX_FILE_BYTES = 5 * 1024 * 1024  # 5 MB
MAX_OFFER_CHARS = 20_000
MAX_DIGEST_CHARS = 1_800

ALLOWED_EXTENSIONS = {".txt", ".md", ".markdown", ".pdf", ".docx"}

SKILL_KEYWORDS = [
    "python", "java", "javascript", "typescript", "go", "golang", "rust", "c++",
    "c#", "ruby", "php", "scala", "kotlin", "swift", "sql", "bash",
    "react", "next.js", "node.js", "django", "fastapi", "spring", "flask",
    "aws", "gcp", "azure", "kubernetes", "docker", "terraform", "linux",
    "postgresql", "postgres", "mysql", "mongodb", "redis", "kafka", "spark",
    "airflow", "dbt", "snowflake", "databricks", "bigquery", "redshift",
    "pandas", "numpy", "pytorch", "tensorflow", "scikit-learn", "mlflow",
    "llm", "rag", "nlp", "computer vision", "machine learning", "deep learning",
    "data engineering", "data science", "mlops", "microservices", "rest",
    "graphql", "grpc", "distributed systems", "system design", "ci/cd",
    "observability", "security", "architecture", "leadership", "mentoring",
]

TRACK_KEYWORDS = {
    "software_engineering": [
        "software engineer", "backend", "frontend", "full-stack", "full stack",
        "api", "microservice", "web application", "mobile", "developer",
    ],
    "architecture": [
        "architect", "architecture", "system design", "distributed system",
        "scalability", "high availability", "microservice",
    ],
    "technical_lead": [
        "tech lead", "technical lead", "engineering manager", "lead engineer",
        "team lead", "mentoring", "line management",
    ],
    "ai_ml": [
        "machine learning", "ml engineer", "ai engineer", "deep learning",
        "llm", "nlp", "computer vision", "mlops", "data scientist",
    ],
    "data": [
        "data engineer", "data analyst", "analytics engineer", "data platform",
        "etl", "elt", "warehouse", "snowflake", "databricks", "airflow",
    ],
}

SENIORITY_PATTERNS = [
    ("principal", r"\bprincipal\b"),
    ("staff", r"\bstaff\b"),
    ("director", r"\bdirector\b"),
    ("head", r"\bhead of\b"),
    ("lead", r"\blead\b"),
    ("senior", r"\bsenior\b|\bsr\.?\b"),
    ("junior", r"\bjunior\b|\bjr\.?\b"),
    ("intern", r"\bintern\b|\binternship\b"),
]


class OfferExtractionError(ValueError):
    """Raised when a job-offer file cannot be read or is unsupported."""


def extract_offer_text(filename: str, content: bytes) -> str:
    """Extract plain text from an uploaded job-offer file."""
    if not content:
        raise OfferExtractionError("The uploaded file is empty.")
    if len(content) > MAX_FILE_BYTES:
        raise OfferExtractionError("The uploaded file is larger than 5 MB.")

    ext = Path(filename or "").suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise OfferExtractionError(
            "Unsupported file type. Upload a .txt, .md, .pdf or .docx file."
        )

    if ext == ".pdf":
        text = _extract_pdf(content)
    elif ext == ".docx":
        text = _extract_docx(content)
    else:
        text = _decode_text(content)

    cleaned = clean_offer(text)
    if not cleaned:
        raise OfferExtractionError(
            "No readable text was found in the file. Try pasting the offer instead."
        )
    return cleaned


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _extract_pdf(content: bytes) -> str:
    if not content.lstrip().startswith(b"%PDF"):
        raise OfferExtractionError("The file does not look like a valid PDF.")
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise OfferExtractionError(
            "PDF support is unavailable on the server. Paste the offer text instead."
        ) from exc

    try:
        reader = PdfReader(io.BytesIO(content))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:
        raise OfferExtractionError(
            "Could not read the PDF. Try pasting the offer text instead."
        ) from exc
    return "\n".join(pages)


def _extract_docx(content: bytes) -> str:
    if not content.startswith(b"PK"):
        raise OfferExtractionError("The file does not look like a valid DOCX.")
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise OfferExtractionError(
            "DOCX support is unavailable on the server. Paste the offer text instead."
        ) from exc

    try:
        document = Document(io.BytesIO(content))
    except Exception as exc:
        raise OfferExtractionError(
            "Could not read the DOCX file. Try pasting the offer text instead."
        ) from exc

    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def clean_offer(text: str) -> str:
    """Normalise whitespace and cap the offer length."""
    cleaned = (text or "").replace("\x00", " ")
    cleaned = re.sub(r"\r\n?", "\n", cleaned)
    cleaned = re.sub(r"[ \t]+", " ", cleaned)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    cleaned = cleaned.strip()
    return cleaned[:MAX_OFFER_CHARS]


def offer_digest(text: str) -> str:
    return clean_offer(text)[:MAX_DIGEST_CHARS]


def heuristic_role_title(text: str) -> str:
    """Best-effort role title extraction when the LLM is unavailable."""
    lines = [line.strip(" -•\t") for line in text.splitlines() if line.strip()]
    for line in lines[:40]:
        match = re.match(
            r"(?:job\s*title|role|position|title)\s*[:\-]\s*(.+)",
            line,
            re.IGNORECASE,
        )
        if match:
            return match.group(1).strip()[:120]

    title_terms = (
        r"\b(?:software|backend|frontend|full[- ]stack|data|machine learning|ml|ai|"
        r"platform|cloud|devops|site reliability|security|solutions|systems|"
        r"enterprise|technical|engineering|staff|principal|senior|lead|junior)\b"
    )
    for line in lines[:15]:
        if 3 <= len(line) <= 120 and re.search(title_terms, line, re.IGNORECASE):
            # Take the role phrase before "at Company", a separator, or a
            # sentence boundary such as ". We need...".
            candidate = re.split(
                r"\s+(?:at|@)\s+|\s*[|•·]\s*|\s+[–-]\s+|\.\s+",
                line,
                maxsplit=1,
            )[0]
            candidate = candidate.strip(" .,:;")
            if 3 <= len(candidate) <= 120:
                return candidate
    return "Software Engineer"


def heuristic_company(text: str) -> str:
    patterns = [
        (
            r"(?:company|employer|organisation|organization)\s*[:\-]\s*([^\n]{2,80})",
            re.IGNORECASE,
        ),
        # Capture up to four capitalised words after "at"/"join", stopping at
        # punctuation so "Acme. We need..." yields "Acme". The keyword is
        # case-insensitive; the captured company name must keep its capitals.
        (
            r"\b(?i:join|at)\s+([A-Z][A-Za-z0-9&.\-]*(?:\s+[A-Z][A-Za-z0-9&.\-]*){0,3})",
            0,
        ),
    ]
    for pattern, flags in patterns:
        match = re.search(pattern, text, flags)
        if match:
            value = match.group(1).strip(" .,-")
            value = re.split(r"\.\s|,", value, maxsplit=1)[0].strip()
            if value and len(value) <= 80:
                return value
    return "the company"


def heuristic_seniority(text: str) -> str:
    lowered = text.lower()
    for name, pattern in SENIORITY_PATTERNS:
        if re.search(pattern, lowered):
            return name
    return "mid-level"


def heuristic_track(text: str) -> str:
    lowered = text.lower()
    title = heuristic_role_title(text).lower()

    # The role title is the strongest signal.
    if re.search(r"\barchitect", title):
        return "architecture"
    if re.search(r"\b(?:data engineer|data analyst|analytics engineer|data platform)", title):
        return "data"
    if re.search(r"\b(?:machine learning|ml|ai|artificial intelligence|data scientist)", title):
        return "ai_ml"
    if re.search(r"\b(?:tech lead|technical lead|lead engineer|engineering manager|head of engineering)", title):
        return "technical_lead"
    if re.search(r"\b(?:software|backend|frontend|full[- ]stack|developer|engineer)", title):
        return "software_engineering"

    scores = {
        track: sum(1 for keyword in keywords if _contains_keyword(keyword, lowered))
        for track, keywords in TRACK_KEYWORDS.items()
    }
    best_track, best_score = max(scores.items(), key=lambda item: item[1])
    return best_track if best_score > 0 else "software_engineering"


def _contains_keyword(keyword: str, lowered: str) -> bool:
    """Match a skill as a word, avoiding false positives such as "go" in Google."""
    return (
        re.search(
            rf"(?<![a-z0-9]){re.escape(keyword)}(?![a-z0-9])",
            lowered,
        )
        is not None
    )


def heuristic_skills(text: str, limit: int = 10) -> list[str]:
    lowered = text.lower()
    unique: list[str] = []
    for skill in SKILL_KEYWORDS:
        if _contains_keyword(skill, lowered) and skill not in unique:
            unique.append(skill)
        if len(unique) >= limit:
            break
    return unique


def heuristic_brief(text: str) -> dict:
    cleaned = clean_offer(text)
    return {
        "role_title": heuristic_role_title(cleaned),
        "company": heuristic_company(cleaned),
        "seniority": heuristic_seniority(cleaned),
        "role_track": heuristic_track(cleaned),
        "key_skills": heuristic_skills(cleaned),
        "offer_digest": offer_digest(cleaned),
    }
