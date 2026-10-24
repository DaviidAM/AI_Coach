"""Tests for job-offer extraction and heuristic fallbacks."""

import io

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.interview.offer import (
    OfferExtractionError,
    clean_offer,
    extract_offer_text,
    heuristic_brief,
    heuristic_company,
    heuristic_role_title,
    heuristic_seniority,
    heuristic_track,
)


def test_extract_txt_and_markdown():
    text = "Senior Backend Engineer\nAcme\nPython, AWS, Kubernetes"
    assert "Senior Backend Engineer" in extract_offer_text("offer.txt", text.encode())
    assert "Acme" in extract_offer_text("offer.md", text.encode())


def test_extract_docx():
    document = Document()
    document.add_heading("Job Title: Data Engineer", level=1)
    document.add_paragraph("Company: DataCorp")
    document.add_paragraph("We need SQL, Spark and Airflow.")
    buffer = io.BytesIO()
    document.save(buffer)

    text = extract_offer_text("offer.docx", buffer.getvalue())
    assert "Data Engineer" in text
    assert "DataCorp" in text
    assert "Airflow" in text


def test_rejects_unsupported_extension():
    with pytest.raises(OfferExtractionError):
        extract_offer_text("offer.exe", b"MZ")


def test_rejects_invalid_pdf():
    with pytest.raises(OfferExtractionError):
        extract_offer_text("offer.pdf", b"not really a pdf")


def test_extract_pdf():
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    page = writer.pages[0]
    stream = DecodedStreamObject()
    stream.set_data(
        b"BT /F1 12 Tf 72 720 Td (Data Engineer at DataCorp) Tj ET"
    )
    page[NameObject("/Contents")] = writer._add_object(stream)
    font = DictionaryObject()
    font.update(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    fonts = DictionaryObject()
    fonts[NameObject("/F1")] = writer._add_object(font)
    resources = DictionaryObject()
    resources[NameObject("/Font")] = fonts
    page[NameObject("/Resources")] = resources

    buffer = io.BytesIO()
    writer.write(buffer)
    text = extract_offer_text("offer.pdf", buffer.getvalue())
    assert "Data Engineer at DataCorp" in text


def test_rejects_empty_text():
    with pytest.raises(OfferExtractionError):
        extract_offer_text("offer.txt", b"   \n\n ")


def test_rejects_oversized_file(monkeypatch):
    from app.interview import offer as offer_module

    monkeypatch.setattr(offer_module, "MAX_FILE_BYTES", 4)
    with pytest.raises(OfferExtractionError):
        extract_offer_text("offer.txt", b"too long")


def test_clean_offer_collapses_whitespace_and_caps_length(monkeypatch):
    from app.interview import offer as offer_module

    monkeypatch.setattr(offer_module, "MAX_OFFER_CHARS", 20)
    cleaned = clean_offer("  Senior   Engineer\r\n\r\n\r\nPython  ")
    assert "\r" not in cleaned
    assert "  " not in cleaned
    assert len(cleaned) <= 20


def test_heuristic_brief_detects_role_company_track_and_skills():
    text = (
        "Senior Backend Engineer at Acme. We need Python, AWS, Kubernetes "
        "and distributed systems experience."
    )
    brief = heuristic_brief(text)
    assert brief["role_title"] == "Senior Backend Engineer"
    assert brief["company"] == "Acme"
    assert brief["seniority"] == "senior"
    assert brief["role_track"] == "software_engineering"
    assert "python" in brief["key_skills"]


def test_heuristic_track_detects_data_ai_architect_and_lead():
    assert heuristic_track("Data Engineer with SQL and Airflow") == "data"
    assert heuristic_track("Machine Learning Engineer building LLM systems") == "ai_ml"
    assert heuristic_track("Software Architect for distributed systems") == "architecture"
    assert heuristic_track("Technical Lead for the platform team") == "technical_lead"


def test_heuristic_skills_uses_word_boundaries():
    from app.interview.offer import heuristic_skills

    assert "go" in heuristic_skills("We build services in Go.")
    assert "go" not in heuristic_skills("We use Google Workspace.")


def test_heuristic_title_company_and_seniority():
    text = "Job Title: Staff Platform Engineer\nCompany: Globex\nJoin Globex."
    assert heuristic_role_title(text) == "Staff Platform Engineer"
    assert heuristic_company(text) == "Globex"
    assert heuristic_seniority(text) == "staff"
